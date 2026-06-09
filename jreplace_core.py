"""
Core functions for jreplace - replacements in Word documents from Excel data.

This module provides reusable functions for replacing keywords in Word documents
with content from Excel spreadsheets. It can be used both by the CLI and by other
applications (e.g., web interfaces).

Author: Carlos Allende Prieto
"""

import os
import shutil
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    """
    Yield each paragraph and table child within *parent*, in document order.
    Each returned value is an instance of either Table or Paragraph. *parent*
    would most commonly be a reference to a main Document object, but
    also works for a _Cell object, which itself can contain paragraphs and tables.
    """
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("something's not right")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def jreplace_file(sheet, template, rownum, tag, output_dir=None):
    """
    Does the replacement job on a single file (template) for a single row 
    in the (excel) sheet.
    
    Args:
        sheet: openpyxl worksheet object
        template: path to the .docx template file
        rownum: row number in the sheet to use for replacement
        tag: string tag to append to output filename (e.g., "_0001")
        output_dir: directory to save output file. If None, uses template directory.
    
    Returns:
        Path to the output file
    """
    # Open word template
    doc = docx.Document(template)
    
    # Find and replace
    for colnum in range(1, sheet.max_column + 1):
        keyword = str(sheet.cell(row=1, column=colnum).value)
        target = str(sheet.cell(row=rownum, column=colnum).value)
        
        if keyword == 'None' or target == 'None':
            continue
        
        keyword = keyword.strip()
        target = target.strip()
        
        for block in iter_block_items(doc):
            if hasattr(block, 'text'):  # deals with page breaks
                if block.text == u'\n':
                    continue
            
            if hasattr(block, 'runs'):  # if it has runs it is a paragraph
                inline = block.runs
                if len(inline) < 1:
                    continue
                for j in range(len(inline)):
                    if inline[j].text == '':
                        continue
                    text = inline[j].text.replace(keyword, target)
                    inline[j].text = text
            else:  # if it doesn't, it is a table
                if not hasattr(block, 'rows'):
                    continue
                for row in block.rows:
                    if not hasattr(row, 'cells'):
                        continue
                    for cell in row.cells:
                        if not hasattr(cell, 'paragraphs'):
                            continue
                        for paragraph in cell.paragraphs:
                            inline = paragraph.runs
                            for j in range(len(inline)):
                                text = inline[j].text.replace(keyword, target)
                                inline[j].text = text
    
    # Save output to a new doc
    if output_dir is None:
        output_dir = os.path.dirname(template)
    
    outdoc = os.path.join(output_dir, os.path.basename(template)[0:-5] + tag + '.docx')
    doc.save(outdoc)
    return outdoc


def getdocxs(directory):
    """
    Gets all the docx files in a directory tree
    """
    files = []
    for directory_path, dirnames, filenames in os.walk(directory):
        for file in filenames:
            fullname = os.path.join(directory_path, file)
            if fullname.lower().endswith('.docx'):
                files.append(fullname)
    return files


def jreplace_dir(sheet, template_dir, rownum, tag, output_dir=None):
    """
    Does the replacement job on all docx documents in a tree dir (template) 
    for a single row in the (excel) sheet.
    
    Args:
        sheet: openpyxl worksheet object
        template_dir: path to directory tree containing templates
        rownum: row number in the sheet to use for replacement
        tag: string tag to append to output directory (e.g., "_0001")
        output_dir: directory to save output folder. If None, uses template_dir parent.
    
    Returns:
        Path to the output directory
    """
    if output_dir is None:
        output_dir = os.path.dirname(template_dir)
    
    newdir = os.path.join(output_dir, os.path.basename(template_dir) + tag)
    shutil.copytree(template_dir, newdir)
    docfiles = getdocxs(newdir)
    
    for file in docfiles:
        jreplace_file(sheet, file, rownum, '', newdir)
    
    return newdir
