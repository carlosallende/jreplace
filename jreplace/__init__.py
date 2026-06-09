"""
jReplace - Replace keywords in Word documents using Excel data

A powerful tool for batch replacing keywords in Word documents using values
from an Excel spreadsheet. Supports both single documents and directory trees
with multiple documents.
"""

__version__ = "1.0.0"
__author__ = "Carlos Allende Prieto"

import sys
import os
import openpyxl
from jreplace.core import jreplace_file, jreplace_dir


def info():
    """Inform users on syntax"""
    print('syntax:  jreplace  xlsx  docx/dirtree  rownumber1 rownumber2 ...')


def prechecks():
    """Carry out basic checks on the command-line arguments"""
    if len(sys.argv) < 4:
        print('ERROR: at least 3 arguments are required')
        info()
        exit()
    if int(sys.argv[3]) < 2:
        print('ERROR: the 3rd argument (rownum) must be >=2')
        info()
        exit()
    if sys.argv[1][-4:] != 'xlsx':
        print(sys.argv[1][-4:])
        print('ERROR: the 1st argument must be the name of the xlsx file')
        info()
        exit()

    if os.path.isfile(sys.argv[2]):
        if sys.argv[2][-4:] != 'docx':
            print('ERROR: the 2nd argument must be the name of the docx file')
            info()
            exit()
    elif os.path.isdir(sys.argv[2]):
        pass
    else:
        print("ERROR: template " + sys.argv[2] + " is neither a file nor a directory")
        info()
        exit()


def main():
    """Main entry point for the jreplace command"""
    # Carry out basic checks
    prechecks()

    # Assign parameters to variables
    workbook = sys.argv[1]
    template = sys.argv[2]
    rownumbers = sys.argv[3:]

    # Open excel sheet
    wb = openpyxl.load_workbook(workbook)
    sheetnames = wb.get_sheet_names()
    sheet = wb.get_sheet_by_name(sheetnames[0])

    # Loop over rownumbers
    for row in rownumbers:
        rownum = int(row)

        if rownum > sheet.max_row:
            print('ERROR: requested rownumber=' + row + ' but the excel sheet has only ' + \
                   str(sheet.max_row) + ' rows')
            continue

        tag = ("_%04d" % rownum)

        if os.path.isfile(template):
            jreplace_file(sheet, template, rownum, tag)
            print('\n-- Writing output to ' + os.path.basename(template)[0:-5] + tag + '.docx')
        elif os.path.isdir(template):
            jreplace_dir(sheet, template, rownum, tag)
            print('\n-- Created output directory: ' + os.path.basename(template) + tag)
        else:
            print("ERROR: template " + template + " is neither a file nor a directory")
            break


if __name__ == "__main__":
    main()
